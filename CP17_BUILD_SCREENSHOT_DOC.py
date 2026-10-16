"""Generate final screenshot Word handoff only from real saved screenshots.

No placeholder or synthetic screenshot is ever substituted. Requires python-docx.
Only run after technical work and screenshots are finalized; user must approve before GitHub.
"""
from __future__ import annotations
import argparse
from pathlib import Path

from CP17_EVIDENCE_AUDIT import REQUIRED, OPTIONAL
DESCRIPTIONS={
'01_full_pytest_suite.png':'The final Project 12 automated regression suite and pass count.',
'02_terraform_validation.png':'Terraform formatting and syntax checks for the AWS infrastructure design.',
'03_mlflow_registered_model.png':'Model p09-attrition-gb version 1, candidate state and model registry lineage.',
'04_docker_container_healthy.png':'The secured Docker API is running and reports healthy.',
'05_secured_prediction_parity.png':'Authenticated requests, unauthorized rejection and 507-record parity. Redact all secrets.',
'06_cp13_drift_report.png':'Offline synthetic drift diagnostics and appropriate result limitations.',
'07_cp16_reliability_audit.png':'Offline failure-injection, recovery, privacy and bounded load checks.',
'08_cp15_ci_configuration.png':'Local CI/release-control verification and GitHub Actions YAML configuration; NOT evidence of a completed GitHub Actions run.',
'09_actual_github_actions_ci.png':'Optional: actual GitHub Actions workflow result after publishing or in an authorized private preview.',
'10_live_https_endpoint.png':'Optional: real HTTPS deployment verified in AWS. Only include if actually deployed.',
'11_live_cloudwatch_dashboard.png':'Optional: genuine live CloudWatch monitoring evidence.',
'12_aws_release_rollback_exercise.png':'Optional: actual approved AWS rollback exercise, not a simulation.',
}

def build(folder:Path, output:Path, include_optional:bool=True)->Path:
    try:
        from docx import Document
        from docx.shared import Inches,Pt
        from docx.image.image import Image as DocxImage
    except ImportError as exc:
        raise RuntimeError('Install python-docx when ready: python -m pip install python-docx==1.2.0') from exc
    missing=[n for n in REQUIRED if not (folder/n).is_file()]
    if missing:
        raise FileNotFoundError('Required final screenshots missing: '+', '.join(missing))
    selected=list(REQUIRED)+([n for n in OPTIONAL if (folder/n).is_file()] if include_optional else [])
    doc=Document()
    sec=doc.sections[0]
    sec.top_margin=Inches(0.70);sec.bottom_margin=Inches(0.65)
    sec.left_margin=Inches(0.85);sec.right_margin=Inches(0.85)
    doc.add_heading('Project 12 | Final Screenshot Handoff',0)
    doc.add_paragraph('Production ML, Cloud Deployment & MLOps Platform')
    doc.add_paragraph('Evidence-only handoff. Generated from the actual PNG files below. The user must confirm these images are saved correctly before final README or GitHub publication.')
    for n in selected:
        image=folder/n
        doc.add_page_break()
        doc.add_heading(n,level=2)
        doc.add_paragraph(DESCRIPTIONS[n])
        original=DocxImage.from_file(str(image))
        ratio=min(6.6/original.width.inches, 6.7/original.height.inches)
        doc.add_picture(str(image), width=Inches(original.width.inches*ratio), height=Inches(original.height.inches*ratio))
        doc.add_paragraph('Source filename: '+n)
    output.parent.mkdir(parents=True,exist_ok=True)
    doc.save(output)
    return output

def main()->int:
    p=argparse.ArgumentParser()
    p.add_argument('--screenshots',type=Path,default=Path('screenshots/final'))
    p.add_argument('--output',type=Path,default=Path('outputs/Project12_Final_Screenshot_Handoff.docx'))
    args=p.parse_args()
    path=build(args.screenshots.resolve(),args.output.resolve())
    print('Created screenshot handoff:',path)
    print('PUBLICATION STILL REQUIRES USER CONFIRMATION')
    return 0

if __name__=='__main__':
    raise SystemExit(main())
